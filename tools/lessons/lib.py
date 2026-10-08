"""折射偵探事件簿 PPT：真實照片＋光光偵探＋進場動畫（飛入／彈出／淡入），按一下揭曉答案。"""
import json, re, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn
from lxml import etree

SW, SH = 13.333, 7.5
FONT = '微軟正黑體'
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
BLANK = prs.slide_layouts[6]

def rgb(h): return RGBColor.from_string(h.lstrip('#'))

def set_alpha(fill_elm_parent, alpha):
    """把 solidFill 加上透明度（alpha 0~100）"""
    sf = fill_elm_parent.find('.//' + qn('a:solidFill'))
    clr = sf[0]
    a = etree.SubElement(clr, qn('a:alpha')); a.set('val', str(int(alpha * 1000)))

def font_run(run, size, color='FFFFFF', bold=True, outline=None, ow=2.5):
    f = run.font; f.size = Pt(size); f.bold = bold; f.color.rgb = rgb(color); f.name = FONT
    rPr = run._r.get_or_add_rPr()
    for tag in ('a:ea', 'a:cs'):
        e = rPr.find(qn(tag))
        if e is None: e = etree.SubElement(rPr, qn(tag))
        e.set('typeface', FONT)
    if outline:
        ln = etree.Element(qn('a:ln')); ln.set('w', str(int(ow * 12700)))
        sf = etree.SubElement(ln, qn('a:solidFill')); c = etree.SubElement(sf, qn('a:srgbClr')); c.set('val', outline)
        rPr.insert(0, ln)

def textbox(sl, x, y, w, h, text, size, color='FFFFFF', align=PP_ALIGN.CENTER, outline=None, ow=2.5, anchor=MSO_ANCHOR.MIDDLE):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, Inches(.05))
    for i, line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align
        r = p.add_run(); r.text = line; font_run(r, size, color, outline=outline, ow=ow)
    return tb

def card(sl, x, y, w, h, text, size, fill, color='1A1A1A', line='FFFFFF', lw=3, shape=MSO_SHAPE.ROUNDED_RECTANGLE, align=PP_ALIGN.CENTER, alpha=None):
    s = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    if alpha is not None: set_alpha(s._element.spPr, alpha)
    if line: s.line.color.rgb = rgb(line); s.line.width = Pt(lw)
    else: s.line.fill.background()
    s.shadow.inherit = False
    tf = s.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ('margin_left', 'margin_right'): setattr(tf, m, Inches(.18))
    for i, line_ in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align
        r = p.add_run(); r.text = line_; font_run(r, size, color)
    try: s.adjustments[0] = 0.25
    except Exception: pass
    return s
def shade(sl, y, h, alpha):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(y), Inches(SW), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = rgb('020810'); set_alpha(s._element.spPr, alpha); s.line.fill.background(); s.shadow.inherit = False
    return s

# ---------- 動畫 XML ----------
class Anim:
    def __init__(self): self.id = 2; self.groups = []   # [('auto'|'click', [(kind, spid, opt)])]
    def nid(self): self.id += 1; return self.id
    def add(self, trig, effects): self.groups.append((trig, effects))

def eff_xml(A, kind, spid, node, delay=0, opt=None):
    opt = opt or {}
    dur = opt.get('dur', 600)
    set_vis = lambda: f'<p:set><p:cBhvr><p:cTn id="{A.nid()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
    def anim(attr, fr, to, d=dur):
        return (f'<p:anim calcmode="lin" valueType="num"><p:cBhvr additive="base"><p:cTn id="{A.nid()}" dur="{d}" fill="hold"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
                f'<p:attrNameLst><p:attrName>{attr}</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav tm="0"><p:val><p:strVal val="{fr}"/></p:val></p:tav>'
                f'<p:tav tm="100000"><p:val><p:strVal val="{to}"/></p:val></p:tav></p:tavLst></p:anim>')
    def fade(trans, d=dur):
        return f'<p:animEffect transition="{trans}" filter="fade"><p:cBhvr><p:cTn id="{A.nid()}" dur="{d}"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
    cid = A.nid()
    if kind == 'fly':
        d = opt.get('dir', 'b'); sub = {'b': 4, 'l': 8, 'r': 2, 't': 1}[d]
        fx = {'l': '0-#ppt_w/2', 'r': '1+#ppt_w/2'}.get(d, '#ppt_x'); fy = {'t': '0-#ppt_h/2', 'b': '1+#ppt_h/2'}.get(d, '#ppt_y')
        body = set_vis() + anim('ppt_x', fx, '#ppt_x') + anim('ppt_y', fy, '#ppt_y'); pid, cls = 2, 'entr'
        head = f'presetID="{pid}" presetClass="{cls}" presetSubtype="{sub}" fill="hold"'
        # 加一點彈性減速
        head += ' decel="100000"'
    elif kind == 'fade':
        body = set_vis() + fade('in'); head = 'presetID="10" presetClass="entr" presetSubtype="0" fill="hold"'
    elif kind == 'zoom':
        body = set_vis() + anim('ppt_w', '0', '#ppt_w') + anim('ppt_h', '0', '#ppt_h') + fade('in')
        head = 'presetID="53" presetClass="entr" presetSubtype="16" fill="hold"'
    elif kind == 'wipe':
        sub = {'l': 8, 'u': 1, 'd': 4, 'r': 2}[opt.get('dir', 'l')]; filt = {'l': 'wipe(left)', 'u': 'wipe(up)', 'd': 'wipe(down)', 'r': 'wipe(right)'}[opt.get('dir', 'l')]
        body = set_vis() + f'<p:animEffect transition="in" filter="{filt}"><p:cBhvr><p:cTn id="{A.nid()}" dur="{dur}"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
        head = f'presetID="22" presetClass="entr" presetSubtype="{sub}" fill="hold"'
    elif kind == 'out':
        body = (fade('out', 400) + f'<p:set><p:cBhvr><p:cTn id="{A.nid()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="399"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
                f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="hidden"/></p:to></p:set>')
        head = 'presetID="10" presetClass="exit" presetSubtype="0" fill="hold"'
    elif kind == 'pulse':  # 強調：放大縮小
        body = (f'<p:animScale><p:cBhvr><p:cTn id="{A.nid()}" dur="400" autoRev="1" fill="hold"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr><p:by x="115000" y="115000"/></p:animScale>')
        head = 'presetID="6" presetClass="emph" presetSubtype="0" fill="hold"'
    return (f'<p:par><p:cTn id="{cid}" {head} grpId="0" nodeType="{node}"><p:stCondLst><p:cond delay="{delay}"/></p:stCondLst>'
            f'<p:childTnLst>{body}</p:childTnLst></p:cTn></p:par>')

def timing_xml(A, txt_ids=()):
    out = []
    for gi, (trig, effects) in enumerate(A.groups):
        oc = A.nid()
        cond = '<p:cond delay="indefinite"/>' + ('<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond>' if trig == 'auto' else '')
        inner = []; t = 0
        if trig == 'auto':  # 一個接一個（每個效果延遲 280ms 開始）
            for k, e in enumerate(effects):
                ic = A.nid(); kind, spid, opt = e
                node = 'afterEffect'
                inner.append(f'<p:par><p:cTn id="{ic}" fill="hold"><p:stCondLst><p:cond delay="{t}"/></p:stCondLst><p:childTnLst>{eff_xml(A, kind, spid, node, 0, opt)}</p:childTnLst></p:cTn></p:par>')
                t += (opt or {}).get('gap', 280)
        else:
            ic = A.nid(); parts = []
            for k, e in enumerate(effects):
                kind, spid, opt = e; parts.append(eff_xml(A, kind, spid, 'clickEffect' if k == 0 else 'withEffect', (opt or {}).get('delay', 0), opt))
            inner.append(f'<p:par><p:cTn id="{ic}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{"".join(parts)}</p:childTnLst></p:cTn></p:par>')
        out.append(f'<p:par><p:cTn id="{oc}" fill="hold"><p:stCondLst>{cond}</p:stCondLst><p:childTnLst>{"".join(inner)}</p:childTnLst></p:cTn></p:par>')
    return ('<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
            '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>' + ''.join(out) +
            '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
            '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>'
            + ('<p:bldLst>' + ''.join(f'<p:bldP spid="{i}" grpId="0" animBg="1"/>' for i in txt_ids) + '</p:bldLst>' if txt_ids else '') + '</p:timing>')

TRANS = ['<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="med"><p:fade/></p:transition>',
         '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="med"><p:push dir="l"/></p:transition>',
         '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="med"><p:cover dir="l"/></p:transition>']

def finish(sl, A, k):
    el = sl._element
    el.append(etree.fromstring(TRANS[k % 3]))
    if A.groups:
        ids = set(sp for _, effs in A.groups for (_, sp, _) in effs)
        txt = [sh.shape_id for sh in sl.shapes if sh.shape_id in ids and sh._element.tag == qn('p:sp') and sh.has_text_frame]
        el.append(etree.fromstring(timing_xml(A, txt)))
def line(sl, x1, y1, x2, y2, color, w, dash=None, arrow=False):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(w)
    if dash: c.line.dash_style = dash
    if arrow:
        ln = c.line._get_or_add_ln(); t = etree.SubElement(ln, qn('a:tailEnd')); t.set('type', 'triangle'); t.set('w', 'med'); t.set('len', 'med')
    return c

def oval(sl, x, y, w, h, color, lw, dash=None, fill=None):
    s = sl.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill: s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    else: s.fill.background()
    if color: s.line.color.rgb = rgb(color); s.line.width = Pt(lw)
    else: s.line.fill.background()
    if dash: s.line.dash_style = dash
    s.shadow.inherit = False
    return s
